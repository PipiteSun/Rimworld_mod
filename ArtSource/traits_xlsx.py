import re, os, sys, glob, tarfile, io
import xml.etree.ElementTree as ET
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MOD = r"E:\SteamLibrary\steamapps\common\RimWorld\Mods\LichPhylactery"
GAME = r"E:\SteamLibrary\steamapps\common\RimWorld"
ISEKAI = r"E:\SteamLibrary\steamapps\workshop\content\294100\3657580708"
ISEKAI_ZH = r"E:\SteamLibrary\steamapps\workshop\content\294100\3659072039"
OUT = sys.argv[1]

def strip(s):
    return re.sub(r"<[^>]+>", "", s or "").strip()

def parse(path):
    try:
        return ET.parse(path).getroot()
    except Exception as e:
        print("skip", path, e); return None

# ---------- 原版 + DLC 特性名（中文） ----------
vanilla = {}
for p in glob.glob(os.path.join(GAME, "Data", "*", "Defs", "**", "*.xml"), recursive=True):
    r = parse(p)
    if r is None: continue
    for t in r.iter("TraitDef"):
        dn = t.findtext("defName")
        if not dn: continue
        labs = [strip(d.findtext("label")) for d in t.iter("li") if d.find("label") is not None]
        vanilla[dn] = {"en": labs or [strip(t.findtext("label"))], "zh": []}
for tp in glob.glob(os.path.join(GAME, "Data", "*", "Languages", "ChineseSimplified*.tar")):
    with tarfile.open(tp) as tf:
        for m in tf.getmembers():
            if "DefInjected/TraitDef" not in m.name or not m.name.endswith(".xml"): continue
            try:
                r = ET.fromstring(tf.extractfile(m).read())
            except Exception: continue
            for el in r:
                mm = re.match(r"(\w+)\.degreeDatas\.([^.]+)\.label$", el.tag) or re.match(r"(\w+)\.label$", el.tag)
                if mm and mm.group(1) in vanilla:
                    vanilla[mm.group(1)]["zh"].append(strip(el.text))

# ---------- Isekai 特性名 ----------
isekai = {}
for p in glob.glob(os.path.join(ISEKAI, "**", "Defs", "**", "*.xml"), recursive=True):
    r = parse(p)
    if r is None: continue
    for t in r.iter("TraitDef"):
        dn = t.findtext("defName")
        if dn:
            labs = [strip(d.findtext("label")) for d in t.iter("li") if d.find("label") is not None]
            isekai[dn] = labs or [strip(t.findtext("label"))]
for p in glob.glob(os.path.join(ISEKAI_ZH, "Languages", "ChineseSimplified", "DefInjected", "TraitDef", "*.xml")):
    r = parse(p)
    if r is None: continue
    zh = defaultdict(list)
    for el in r:
        mm = re.match(r"(\w+)\.degreeDatas\.([^.]+)\.label$", el.tag) or re.match(r"(\w+)\.label$", el.tag)
        if mm: zh[mm.group(1)].append(strip(el.text))
    for k, v in zh.items():
        if k in isekai: isekai[k] = v

def other_name(dn):
    if dn in vanilla:
        v = vanilla[dn]; return "原版：" + "/".join(v["zh"] or v["en"])
    if dn in isekai:
        return "Isekai：" + "/".join(isekai[dn])
    return dn

# ---------- 本模组特性 ----------
CATS = {
    "Traits_Backgrounds.xml": "背景特性",
    "HeroTraits.xml": "强力特性",
    "Adventurer_Talents.xml": "稀有天赋",
    "DndTraits.xml": "性格特性 / 诅咒之力",
    "CombatGifts.xml": "武艺天赋",
    "Faith.xml": "信仰",
    "Ascension.xml": "成神",
    "Twilight.xml": "暮光",
}
TAGNAMES = {
    "LP_Background": "背景特性名额（每人最多 1 个）",
    "LP_Curse": "诅咒之力之间互斥（每人最多 1 个诅咒）",
    "LP_Destiny": "天命之间互斥（每人最多 1 个天命）",
    "LP_PointTalent": "稀有天赋之间互斥（每人最多 1 个）",
}
abstract = {}
traits = []
files = sorted(glob.glob(os.path.join(MOD, "Defs", "**", "*.xml"), recursive=True) + glob.glob(os.path.join(MOD, "Bridges", "*", "Defs", "**", "*.xml"), recursive=True))
for p in files:
    r = parse(p)
    if r is None: continue
    for t in r.findall("TraitDef"):
        if t.get("Name"): abstract[t.get("Name")] = t
for p in files:
    r = parse(p)
    if r is None: continue
    for t in r.findall("TraitDef"):
        if t.get("Abstract") == "True": continue
        dn = t.findtext("defName")
        if not dn: continue
        chain = [t]; par = t.get("ParentName")
        while par and par in abstract:
            chain.append(abstract[par]); par = abstract[par].get("ParentName")
        def first(tag):
            for c in chain:
                v = c.find(tag)
                if v is not None: return v
            return None
        tags = []; confl = []
        for c in chain:
            e = c.find("exclusionTags")
            if e is not None: tags += [li.text.strip() for li in e.findall("li") if li.text]
            e = c.find("conflictingTraits")
            if e is not None: confl += [li.text.strip() for li in e.findall("li") if li.text]
        degs = []
        dd = first("degreeDatas")
        if dd is not None:
            for li in dd.findall("li"):
                degs.append((li.findtext("degree") or "0", strip(li.findtext("label")), strip(li.findtext("description"))))
        else:
            degs.append(("0", strip(t.findtext("label")), strip(t.findtext("description"))))
        com = first("commonality"); comf = first("commonalityFemale")
        rx = None
        for c in chain:
            me = c.find("modExtensions")
            if me is None: continue
            for li in me.findall("li"):
                if li.get("Class", "").endswith("TraitRarityExt"): rx = li
            if rx is not None: break
        rar = []
        for d in degs:
            tier = rx.findtext("tier") if rx is not None else None; mark = rx.findtext("mark") if rx is not None else None
            if rx is not None and rx.find("degrees") is not None:
                for o in rx.find("degrees").findall("li"):
                    a = o.text.split(":")
                    if a[0] == d[0]:
                        tier = a[1]
                        if len(a) > 2: mark = None if a[2] == "None" else a[2]
            rar.append((tier, mark, rx.findtext("kind") if rx is not None else ""))
        traits.append({
            "def": dn, "cat": CATS.get(os.path.basename(p), os.path.basename(p)),
            "file": os.path.relpath(p, MOD), "tags": sorted(set(tags)), "confl": confl,
            "degs": degs, "com": com.text if com is not None else "（默认 1）", "rar": rar,
        })
mine = {t["def"]: t for t in traits}

def name(dn):
    if dn in mine:
        return "/".join(d[1] for d in mine[dn]["degs"])
    return other_name(dn)

# Isekai 联动补丁（只在装了 Isekai 时生效）
isekai_extra = defaultdict(set)
pt = open(os.path.join(MOD, "Bridges", "Isekai", "Patches", "IsekaiHeroTraits.xml"), encoding="utf-8").read()
for m in re.finditer(r'<match Class="PatchOperationAdd"><xpath>Defs/TraitDef\[defName="(\w+)"\]/conflictingTraits</xpath><value>(.*?)</value>', pt):
    for li in re.findall(r"<li>(\w+)</li>", m.group(2)):
        isekai_extra[m.group(1)].add(li); isekai_extra[li].add(m.group(1))

# ---------- 冲突对 ----------
pairs = {}
def addpair(a, b, why):
    k = tuple(sorted([a, b]))
    pairs.setdefault(k, set()).add(why)
for t in traits:
    for c in t["confl"]:
        addpair(t["def"], c, "定义里写明互斥（conflictingTraits）")
for a, bs in isekai_extra.items():
    for b in bs:
        if a in mine or b in mine:
            addpair(a, b, "装了 Isekai 时互斥（联动补丁）")
def deg(dn, d):
    for x in mine.get(dn, {"degs": []})["degs"]:
        if x[0] == str(d): return x[1]
    return dn + "(" + str(d) + ")"
code_rules = [
    ("LP_GiftedScore", "Isekai 力大 / 敏捷 / 坚韧 / 聪慧 / 开悟 / 巧舌", "装了 Isekai 时，天赋·某项和对应属性的 Isekai 白色属性特性互斥（代码按等级判断，HeroTraitsIsekai.cs）"),
    (deg("LP_Morality", 1), other_name("Psychopath") + "、" + other_name("Bloodlust") + "、" + deg("LP_UndeadStance", 1), "代码互斥（DndTraits.Clash，按等级）"),
    (deg("LP_Morality", -1), other_name("Kind") + "、" + deg("LP_UndeadStance", -1), "代码互斥（DndTraits.Clash，按等级）"),
    ("/".join(d[1] for d in mine["LP_WildMagic"]["degs"]) if "LP_WildMagic" in mine else "LP_WildMagic", name("LP_ArcaneGift") + "、" + name("LP_Magephobe"), "代码互斥（DndTraits.Clash）"),
    (name("LP_Order"), name("LP_Morality"), "阵营类每人最多 1 个（性格特性生成时）"),
]
groups = defaultdict(list)
for t in traits:
    for g in t["tags"]:
        groups[g].append(t["def"])

# ---------- 稀有度（2026-10-29） ----------
TIER = {"Common": ("普通", None), "Uncommon": ("优秀", "7FBF6A"), "Rare": ("稀有", "5B9BD5"), "Epic": ("史诗", "B070E0"), "Legendary": ("传说", "E08A3A"), "Mythic": ("神话", "E0B040")}
MARKS = {"Curse": ("诅咒", "D05050"), "Deed": ("事迹", "50C0C0")}
TORDER = ["Mythic", "Legendary", "Epic", "Rare", "Uncommon", "Common"]
def rar_text(r):
    tier, mark, kind = r
    if tier is None: return "—"
    return (MARKS[mark][0] + " · " if mark else "") + TIER[tier][0]
def rar_color(r):
    tier, mark, kind = r
    if tier is None: return None
    return MARKS[mark][1] if mark else TIER[tier][1]

# ---------- 写 Excel ----------
wb = Workbook()
hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="4F6D8F")
wrap = Alignment(wrap_text=True, vertical="top")
def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]: c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = wrap
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

ws = wb.active; ws.title = "特性一览"
rows = []
order = list(CATS.values())
for t in sorted(traits, key=lambda x: (order.index(x["cat"]) if x["cat"] in order else 99, x["def"])):
    labels = " / ".join(d[1] for d in t["degs"])
    groupdesc = "；".join(TAGNAMES.get(g, "互斥组 " + g) + "（同组 " + str(len(groups[g]) - 1) + " 个）" for g in t["tags"])
    direct = "、".join(name(c) for c in t["confl"])
    isk = "、".join(name(c) for c in sorted(isekai_extra.get(t["def"], [])) if c not in mine)
    desc = "\n".join((("【" + d[1] + "】") if len(t["degs"]) > 1 else "") + d[2][:160] for d in t["degs"])
    rt = sorted(set(rar_text(r) for r in t["rar"]), key=lambda x: x)
    rows.append([t["cat"], rar_text(t["rar"][0]) if len(rt) == 1 else " / ".join(rar_text(r) for r in t["rar"]), labels, t["def"], t["com"], groupdesc or "无", direct or "无", isk or "", desc])
sheet(ws, ["类别", "稀有度", "名称（各等级）", "defName", "随机权重 commonality", "同名额 / 互斥组", "直接互斥的特性", "装 Isekai 时另外互斥", "说明（截取）"], rows, [14, 12, 26, 26, 12, 40, 40, 36, 80])
for row, t in zip(ws.iter_rows(min_row=2), sorted(traits, key=lambda x: (order.index(x["cat"]) if x["cat"] in order else 99, x["def"]))):
    c = rar_color(t["rar"][0])
    if c:
        row[1].font = Font(bold=True, color="FFFFFF"); row[1].fill = PatternFill("solid", fgColor=c); row[2].font = Font(bold=True, color=c)

wsr = wb.create_sheet("稀有度", 1)
rrow = []
for k in TORDER:
    for t in traits:
        for d, r in zip(t["degs"], t["rar"]):
            if r[0] == k: rrow.append([TIER[k][0] + ("（" + MARKS[r[1]][0] + "）" if r[1] else ""), d[1], t["def"] + ("" if len(t["degs"]) == 1 else "(" + d[0] + ")"), r[2], d[2][:120]])
sheet(wsr, ["稀有度", "名称", "defName(程度)", "类别", "说明（截取）"], rrow, [14, 22, 30, 14, 90])
for row in wsr.iter_rows(min_row=2):
    tn = row[0].value; base = [k for k in TIER if TIER[k][0] == tn.split("（")[0]][0]
    c = MARKS["Curse"][1] if "诅咒" in tn else MARKS["Deed"][1] if "事迹" in tn else TIER[base][1]
    if c: row[0].fill = PatternFill("solid", fgColor=c); row[0].font = Font(bold=True, color="FFFFFF"); row[1].font = Font(bold=True, color=c)

ws2 = wb.create_sheet("冲突对照")
prow = []
for (a, b), why in sorted(pairs.items(), key=lambda kv: (name(kv[0][0]), name(kv[0][1]))):
    prow.append([name(a), a, name(b), b, "；".join(sorted(why))])
for a, b, why in code_rules:
    prow.append([name(a) if a in mine else a, a if a in mine else "", b, "", why])
sheet(ws2, ["特性 A", "defName A", "特性 B", "defName B", "原因"], prow, [26, 24, 30, 24, 50])

ws3 = wb.create_sheet("互斥组")
grow = []
for g, members in sorted(groups.items(), key=lambda kv: -len(kv[1])):
    grow.append([TAGNAMES.get(g, g), g, len(members), "、".join(name(m) for m in sorted(members, key=name))])
sheet(ws3, ["互斥组 / 名额", "exclusionTag", "成员数", "成员（组内任意两个不能同时拥有）"], grow, [40, 22, 8, 120])

ws4 = wb.create_sheet("获得规则")
rules = [
    ["背景特性", "进原版特性随机池，新生成的人形随机；出现权重 0.3 × 设置“背景特性出现率”；部分 DND 背景故事按概率带上对应背景特性", "每人最多 1 个（LP_Background）", "贵族 × 原版苦行"],
    ["稀有度与颜色", "普通 白 / 优秀 绿 / 稀有 蓝 / 史诗 紫 / 传说 橙 / 神话 金；诅咒之力 红、传说事迹 青（档位写在悬停说明第一行）。颜色只表示强度，不改出现概率", "", "详见“稀有度”页"],
    ["强力特性", "新生成时随机一次：普通种族 英雄之资系与天命（传说 / 神话）1% / 加点与强化天赋（稀有，反骨、勤勉史诗）4%，中等种族 0.4% / 3%，传奇种族 只有后者 1%（只抽反骨、勤勉和六个强化天赋）；× 设置出现率；得到一个之后每次 25%（设置）再得一个，按同样比例；剧情首领、宿敌冠军、袭击领队会带", "强力特性 / 诅咒 / 稀有天赋合计每人最多 3 个（设置 1～6）；天命之间互斥；满了再得会取代最早的", "装 Isekai 时与 Isekai 主角类特性互斥（见冲突对照）"],
    ["稀有天赋（天赋异禀、大器晚成、百炼成钢、凡人之志，史诗）", "并入强力特性的同一次随机，2%（传奇种族没有）", "算强力特性数量；四个天赋之间互斥", "种族亲和与天赋作用于同一项时取较大，不相乘"],
    ["诅咒之力（红）", "普通种族 1.5% / 中等 1% / 传奇 0.5%（× 强力特性出现率）；首领带强力特性时 15% 是红色", "算强力特性数量；每人最多 1 个诅咒", ""],
    ["传说事迹（青）", "不随机，由事件授予（屠龙、斩巨人等）", "—", ""],
    ["性格特性", "只用原版 3 个特性名额里的空位：有空位 40% 带 1 个，再有空位 15% 带第 2 个（× 性格特性出现率）；派系 / 种族加权", "每人最多 2 个，阵营类最多 1 个", "不会挤掉原版特性"],
    ["武艺天赋（稀有）", "开局殖民者 35%，普通人形 8%（传奇种族 4%），首领 / 宿敌冠军 50%；击杀 100 / 300 人形、DND 里程碑各 25% 觉醒；巨灵许愿；老存档补一次 35%；得到一个之后每次 25%（设置）再得一个不同的", "和强力特性分开计数，每人最多 3 个（设置 1～13）；满了再得会取代最早的", ""],
    ["信徒 / 无神者", "随机权重 信徒 0.3、无神者 0.1；关闭信仰系统后不再出现", "两者互斥", "和背景特性分开算"],
    ["共同", "只在新生成的人形上随机（13 岁生日补一次）；死灵造物、动物、机械体没有；读旧存档不补发（武艺天赋除外）", "", ""],
]
sheet(ws4, ["类别", "怎么获得", "名额", "其他冲突"], rules, [26, 90, 36, 40])

wb.save(OUT)
print("traits", len(traits), "pairs", len(pairs), "groups", len(groups))
